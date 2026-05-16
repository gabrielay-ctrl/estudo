from docx import Document
from docx.shared import Pt
import os

doc = Document()
doc.add_heading('Entrega Inicial - Bootcamp', 0)

# Informações Pessoais
info = [
    "Nome: Gabriela Yasmin Conceição Viana | RA: 22505273",
    "Professor: Dr. Romes Heriberto",
    "Disciplina: Bootcamp",
    "Curso: Engenharia de Software",
]

for linha in info:
    doc.add_paragraph(linha)

doc.add_heading('Informações do Projeto', level=1)

p1 = doc.add_paragraph()
p1.add_run("Nome do projeto: ").bold = True
p1.add_run("MindStep Tracker")

p2 = doc.add_paragraph()
p2.add_run("Link do repositório público no GitHub: ").bold = True
p2.add_run("https://github.com/gabrielay-ctrl/estudo")

p3 = doc.add_paragraph()
p3.add_run("Breve descrição da proposta:").bold = True

resumo = (
    "O MindStep Tracker é uma aplicação de linha de comando (CLI) voltada para auxiliar pessoas neurodivergentes "
    "(como portadores de TDAH e ansiedade crônica) que sofrem com paralisia de tarefas e esquecimento de "
    "autocuidado básico.\n\nO sistema resolve isso forçando a divisão de grandes atividades cotidianas "
    "em no mínimo 3 \"micro-passos\" antes de sua execução, proporcionando acompanhamento guiado e "
    "micro-recompensas para evitar a sobrecarga cognitiva e garantir o progresso passo a passo."
)

doc.add_paragraph(resumo)

# Salvar o documento Word na Área de Trabalho
file_path = r"C:\Users\gabit\OneDrive\Área de Trabalho\Entrega_Bootcamp_Gabriela.docx"
try:
    doc.save(file_path)
    print(f"Word gerado com sucesso em: {file_path}")
except Exception as e:
    alt_path = r"C:\Users\gabit\.gemini\antigravity\scratch\estudo\Entrega_Bootcamp_Gabriela.docx"
    doc.save(alt_path)
    print(f"Salvo no caminho alternativo: {alt_path}")
